"""Pruebas de las verificaciones de concreto, sin SAP.

Cubre lo que no toca `test_reference.py` (que valida el modelo y el .s2k):

  * diseno por torsion segun ACI 318 22.7, contrastado a mano
  * flexion cuando la seccion no da de si
  * franjas de diseno de la pantalla y momento en la cara del apoyo
  * casos modal y P-Delta en el .s2k
  * generacion del reporte Excel completo

    python tests/test_design.py
"""
from __future__ import annotations

import math
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))

from app.core.builder import G_MURO, G_PILAS, G_VIGA, build
from app.exporters import s2k
from app.results import design as D
from app.results import extract, report
from test_reference import reference_project

FAILS: list[str] = []


def check(label: str, got, want, tol: float = 0.005) -> None:
    if isinstance(want, (int, float)) and isinstance(got, (int, float)) and not isinstance(want, bool):
        err = abs(got - want) / abs(want) * 100.0 if want else abs(got)
        ok = err <= tol * 100.0
        print("  %-46s %12.3f  esperado %12.3f  %s"
              % (label, got, want, "OK" if ok else "<<< FALLA"), flush=True)
    else:
        ok = got == want
        print("  %-46s %-12s  esperado %-12s  %s"
              % (label, got, want, "OK" if ok else "<<< FALLA"), flush=True)
    if not ok:
        FAILS.append(label)


# --------------------------------------------------------------------------
def test_torsion() -> None:
    """Contraste a mano de ACI 318-19 22.7 para la viga de 0.80 x 1.00 m."""
    print("\n1. Diseno por torsion (viga 0.80 x 1.00, f'c 21 MPa, fy 420 MPa)")
    b, h, d, cover = 0.80, 1.00, 0.942, 0.05
    fc, fy = 21000.0, 420000.0
    Tu, Vu = 385.0, 300.0
    t = D.torsion_design(Tu, Vu, b, h, d, fc, fy, cover, tie_size="#4", long_bar_size="#5")

    fc_m = 21.0
    acp, pcp = 800.0 * 1000.0, 2.0 * (800.0 + 1000.0)
    check("umbral 22.7.4 [kN-m]", t.T_threshold,
          0.75 * 0.083 * math.sqrt(fc_m) * acp ** 2 / pcp / 1e6)
    check("Tcr 22.7.5 [kN-m]", t.T_cr, 0.33 * math.sqrt(fc_m) * acp ** 2 / pcp / 1e6)
    check("torsion significativa", t.significant, True)

    db = D.bar_diam("#4")
    x1, y1 = 800.0 - 2 * 50.0 - db, 1000.0 - 2 * 50.0 - db
    check("Aoh [mm2]", t.Aoh, x1 * y1, 0.001)
    check("ph [mm]", t.ph, 2.0 * (x1 + y1), 0.001)
    check("Ao = 0.85 Aoh [mm2]", t.Ao, 0.85 * x1 * y1, 0.001)

    At_s = Tu * 1e6 / (0.75 * 2.0 * 0.85 * x1 * y1 * 420.0)      # mm2/mm por rama
    check("At/s por rama [mm2/m]", t.At_s, At_s * 1000.0, 0.01)
    check("Al longitudinal [mm2]", t.Al, At_s * 2.0 * (x1 + y1), 0.01)
    check("(Av+2At)/s [mm2/m]", t.Avt_s, 2.0 * At_s * 1000.0, 0.01)
    check("s maxima 9.7.6.3.3 [mm]", t.s_max, min(2.0 * (x1 + y1) / 8.0, 300.0), 0.001)

    Vc = 0.17 * math.sqrt(fc_m) * 800.0 * 942.0
    dem = math.sqrt((300000.0 / (800.0 * 942.0)) ** 2
                    + (385e6 * 2.0 * (x1 + y1) / (1.7 * (x1 * y1) ** 2)) ** 2)
    cap = 0.75 * (Vc / (800.0 * 942.0) + 0.66 * math.sqrt(fc_m))
    check("D/C limite de seccion 22.7.7.1", t.stress_ratio, dem / cap, 0.01)
    check("seccion admisible", t.section_ok, True)

    print("\n   Torsion despreciable y redistribucion de compatibilidad")
    low = D.torsion_design(20.0, 300.0, b, h, d, fc, fy, cover)
    check("Tu bajo el umbral -> no significativa", low.significant, False)
    check("Tu bajo el umbral -> At/s nulo", low.At_s, 0.0, 1.0)
    comp = D.torsion_design(385.0, 300.0, b, h, d, fc, fy, cover, compatibility=True)
    check("compatibilidad -> redistribuye", comp.redistributed, True)
    check("compatibilidad -> At/s menor", comp.At_s < t.At_s, True)


# --------------------------------------------------------------------------
def test_flexion() -> None:
    """Cuando la seccion no da de si, el D/C debe medir el deficit real."""
    print("\n2. Flexion con seccion insuficiente")
    kw = dict(b=0.8, h=1.0, fc=21000.0, fy=420000.0, cover=0.05)
    ok_case = D.flexure_rect(Mu=300.0, **kw)
    check("caso normal: ok", ok_case.ok, True)
    check("caso normal: As_req < As_max", ok_case.As_req < ok_case.As_max, True)

    prev = None
    for Mu in (3200.0, 3600.0, 5000.0):
        bad = D.flexure_rect(Mu=Mu, **kw)
        check("Mu=%.0f: no cumple" % Mu, bad.ok, False)
        check("Mu=%.0f: limitada por seccion" % Mu, bad.limitada_por_seccion, True)
        check("Mu=%.0f: As_req topado en As_max" % Mu, bad.As_req, bad.As_max, 0.001)
        check("Mu=%.0f: D/C > 1" % Mu, bad.ratio > 1.0, True)
        if prev is not None:
            check("Mu=%.0f: D/C crece con Mu" % Mu, bad.ratio > prev, True)
        prev = bad.ratio


# --------------------------------------------------------------------------
def _synthetic_rows(p, m):
    """Fuerzas plausibles como las devolveria SAP, para ejercitar el post-proceso."""
    pile_ys = p.pile_y_positions()
    L = p.piles.spacing
    cases = (("U1 1.2D+1.6H", 1.0), ("U2 1.2D+1.6H+1.0S", 1.25))

    wall = []
    for aid in m.group_element_ids(G_MURO, "Area"):
        for jid in m.areas[aid - 1].joints:
            j = m.get_joint(jid)
            zt, zb = p.wall.crown_at(j.y), p.wall.base_at(j.y)
            hh = max(zt - zb, 1e-6)
            u = min(min(abs(j.y - yp) for yp in pile_ys) / (L / 2.0), 1.0)
            m22 = -160.0 * ((zt - j.z) / hh) ** 3
            m11 = -90.0 * (1 - 3 * u * u) * ((zt - j.z) / hh) ** 2
            v = 60.0 * ((zt - j.z) / hh) ** 2
            for case, k in cases:
                wall.append({"obj": aid, "point": jid, "case": case,
                             "F11": 0.0, "F22": 0.0, "F12": 0.0,
                             "M11": m11 * k, "M22": m22 * k, "M12": 0.0,
                             "V13": 0.0, "V23": v * k, "VMax": v * k})

    beam = []
    for fid in m.group_element_ids(G_VIGA, "Frame"):
        fr = m.frames[fid - 1]
        ym = 0.5 * (m.get_joint(fr.i).y + m.get_joint(fr.j).y)
        u = min(min(abs(ym - yp) for yp in pile_ys) / (L / 2.0), 1.0)
        for case, k in cases:
            beam.append({"obj": fid, "sta": 0.0, "case": case, "step": "",
                         "P": -50.0 * k, "V2": 0.0, "V3": 300.0 * (1 - u) * k,
                         "T": 385.0 * k, "M2": 420.0 * (1 - 3 * u * u) * k, "M3": 0.0})

    pile = []
    for fid in m.group_element_ids(G_PILAS, "Frame"):
        fr = m.frames[fid - 1]
        zm = 0.5 * (m.get_joint(fr.i).z + m.get_joint(fr.j).z)
        f = max(0.0, 1.0 - abs(zm) / 10.0)
        for case, k in cases:
            pile.append({"obj": fid, "sta": 0.0, "case": case, "step": "",
                         "P": -800.0 * k, "V2": 120.0 * f * k, "V3": 0.0,
                         "T": 5.0 * k, "M2": 260.0 * f * k, "M3": 0.0})
    return pile, beam, wall


def test_franjas() -> None:
    print("\n3. Franjas de diseno de la pantalla")
    p = reference_project()
    p.mesh.wall_design_strips = 3
    res = build(p)
    m = res.model
    pile_rows, beam_rows, wall_rows = _synthetic_rows(p, m)
    checks = extract.verify(p, m, pile_rows, beam_rows, wall_rows)

    fr = checks["muro"].get("franjas")
    check("hay franjas", bool(fr), True)
    check("numero de franjas", fr["n_franjas"], 3, 0.0)
    check("filas en altura", fr["n_filas"], 6, 0.0)
    check("cota de la cara de la viga [m]", fr["z_cara_viga"],
          p.wall.z_base + p.cap_beam.depth / 2.0, 0.001)

    f1, f2, f3 = fr["franjas"]
    check("M22 decrece hacia la corona", f1["M22"] > f2["M22"] > f3["M22"], True)
    check("As vertical no crece hacia la corona",
          f1["vertical"]["As"] >= f2["vertical"]["As"] >= f3["vertical"]["As"], True)
    check("franjas cubren toda la altura", f3["z_fin"] - f1["z_ini"],
          p.wall.z_top - p.wall.z_base, 0.001)
    check("momento en la cara <= momento en el eje",
          all(f["M11_cara"] <= f["M11_eje"] + 1e-9 for f in fr["franjas"]), True)
    check("la cara de la pila reduce el momento", fr["reduccion_cara_horizontal_pct"] > 0.0, True)
    check("As maximo por franja <= As del maximo global",
          fr["As_vertical_max"] <= max(checks["muro"]["vertical"]["As"],
                                       checks["muro"]["horizontal"]["As"]) + 1e-6, True)

    print("\n   %-7s %-13s %9s %9s %9s %9s %9s %9s"
          % ("franja", "z [m]", "M22", "M22 dis", "As vert", "M11 eje", "M11 cara", "As h ap"))
    for f in fr["franjas"]:
        print("   %-7d %5.2f-%-6.2f %9.1f %9.1f %9.0f %9.1f %9.1f %9.0f"
              % (f["franja"], f["z_ini"], f["z_fin"], f["M22"], f["M22_diseno"],
                 f["vertical"]["As"], f["M11_eje"], f["M11_cara"],
                 f["horizontal_apoyo"]["As"]))
    return p, res, checks


# --------------------------------------------------------------------------
def test_datos_para_la_interfaz(checks) -> None:
    """Lo que la pagina necesita para dibujar: curvas, envolventes y mapa."""
    print("\n3b. Datos expuestos para la interfaz")
    inter = checks.get("interaccion") or {}
    check("hay curva de interaccion por zona", len(inter) >= 1, True)
    for nombre, z in inter.items():
        check("curva de %s con puntos" % nombre, len(z["curva"]) > 10, True)
        check("curva de %s ordenada en P" % nombre,
              all(a[0] <= b[0] + 1e-6 for a, b in zip(z["curva"], z["curva"][1:])), True)
        check("demanda de %s no vacia" % nombre, len(z["demanda"]) > 0, True)
        check("compresion positiva en %s" % nombre, z["curva"][-1][0] > 0, True)
        check("traccion pura en %s" % nombre, z["curva"][0][0] < 0, True)

    env = checks.get("envolventes") or {}
    check("envolventes de pilas", len(env.get("pilas") or []) > 0, True)
    check("envolventes de viga", len(env.get("viga_cabezal") or []) > 0, True)
    una = (env.get("viga_cabezal") or [{}])[0]
    for comp in ("P", "V2", "V3", "T", "M2", "M3"):
        check("la envolvente trae %s con su combinacion" % comp,
              bool(una.get(comp, {}).get("caso")), True)

    els = (checks.get("muro") or {}).get("elementos") or []
    check("valor por shell del muro", len(els), 90, 0.0)
    check("cada shell trae M11, M22 y V",
          all(all(k in e for k in ("M11", "M22", "V")) for e in els), True)


def test_torsion_en_viga(checks) -> None:
    print("\n4. Torsion de la viga cabezal dentro de la verificacion")
    tv = checks.get("torsion_viga")
    check("hay diseno de torsion", bool(tv), True)
    if not tv:
        return
    check("Tu supera el umbral", tv["Tu"] > tv["T_umbral"], True)
    check("At/s positivo", tv["At_s"] > 0.0, True)
    check("Al positivo", tv["Al"] > 0.0, True)
    check("separacion <= s maxima", tv["s_estribo"] <= tv["s_max"], True)
    check("al menos 4 barras en el perimetro", tv["n_barras_torsion"] >= 4, True)
    print("   Tu = %.1f kN-m (umbral %.1f) -> estribos %s @ %.0f mm, %d barras %s, Al = %.0f mm2"
          % (tv["Tu"], tv["T_umbral"], tv["estribo"], tv["s_estribo"],
             tv["n_barras_torsion"], tv["barra_torsion"], tv["Al"]))


# --------------------------------------------------------------------------
def test_casos_analisis() -> None:
    print("\n5. Casos modal y P-Delta en el .s2k")
    p = reference_project()
    base = s2k.export(build(p).model)
    check("sin opciones no hay tabla modal", "CASE - MODAL 1" in base, False)
    check("sin opciones no hay tabla no lineal", "NONLINEAR PARAMETERS" in base, False)

    p.analysis.modal = True
    p.analysis.modal_modes = 10
    p.analysis.pdelta = True
    res = build(p)
    txt = s2k.export(res.model)
    check("caso modal definido", "Case=MODAL   Type=LinModal" in txt, True)
    check("tabla modal presente", "CASE - MODAL 1 - GENERAL" in txt, True)
    check("numero de modos", "MaxNumModes=10" in txt, True)
    check("caso no lineal definido", "Case=PDELTA   Type=NonStatic" in txt, True)
    check("P-Delta activado", "GeoNonLin=P-Delta" in txt, True)
    check("P-Delta lleva sus dos patrones",
          txt.count("Case=PDELTA   LoadType=") , 2, 0.0)
    check("termina bien", txt.rstrip().endswith("END TABLE DATA"), True)


# --------------------------------------------------------------------------
def test_reporte(p, res, checks) -> None:
    print("\n6. Reporte Excel con todas las hojas")
    checks = dict(checks)
    checks["modal"] = {
        "caso": "MODAL", "n_modos": 3, "T1": 0.412,
        "modo_dominante_X": 1, "T_dominante_X": 0.412,
        "modo_dominante_Y": 2, "T_dominante_Y": 0.233,
        "masa_acumulada_X": 0.91, "masa_acumulada_Y": 0.88,
        "modos": [{"modo": i, "T": 0.412 / i, "f": i / 0.412,
                   "Ux": 0.6 / i, "Uy": 0.2 * i, "Uz": 0.0,
                   "SumUx": 0.6, "SumUy": 0.4, "SumUz": 0.0} for i in (1, 2, 3)],
        "nota": "",
    }
    checks["diseno_sap"] = {
        "codigo": "ACI 318-19",
        "viga_cabezal": [{"frame": "1", "x": 0.0, "As_sup": 2100.0, "As_inf": 1800.0,
                          "Av_s": 900.0, "Al_torsion": 3500.0, "At_s_torsion": 1150.0,
                          "error": "", "aviso": ""}],
        "pilas": [{"frame": "10", "x": 0.0, "As_PMM": 6400.0, "ratio_PMM": 0.71,
                   "Av_s_mayor": 700.0, "error": "", "aviso": ""}],
    }
    checks["comparacion_torsion"] = extract._compare_torsion(
        checks.get("torsion_viga"), checks["diseno_sap"])
    check("comparacion de torsion generada", bool(checks["comparacion_torsion"]), True)

    out = os.path.join(ROOT, "tests", "_out_diseno.xlsx")
    report.save_report(p, res.report, out, checks, res.warnings)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    print("   hojas:", ", ".join(wb.sheetnames))
    for hoja in ("Datos", "Pilas", "VigaCabezal", "Muro", "Torsion",
                 "MuroFranjas", "Modal", "DisenoSAP"):
        check("hoja %s" % hoja, hoja in wb.sheetnames, True)
    check("archivo escrito", os.path.getsize(out) > 5000, True)



# --------------------------------------------------------------------------
def test_apoyos_de_pila() -> None:
    """Solo la punta lleva restriccion; el fuste, solo balasto."""
    print("\n7. Apoyos de las pilas")
    from app.core.params import WallProject

    p = WallProject()
    m = build(p).model
    n_pilas = len(p.pile_y_positions())

    check("una restriccion por pila", len(m.restraints), n_pilas, 0.0)
    cotas = {round(m.get_joint(r.joint).z, 3) for r in m.restraints}
    check("todas en la cota de punta", cotas, {round(p.piles.z_bot, 3)})
    gdl = {(r.u1, r.u2, r.u3, r.r1, r.r2, r.r3) for r in m.restraints}
    check("solo U3 restringido", gdl, {(False, False, True, False, False, False)})

    restringidos = {r.joint for r in m.restraints}
    con_resorte = {s.joint for s in m.springs}
    check("los nudos del fuste no llevan restriccion",
          len(con_resorte - restringidos), len(con_resorte) - n_pilas, 0.0)
    check("ningun resorte carga en vertical",
          sum(1 for s in m.springs if s.u3 > 0), 0, 0.0)
    check("todos los nudos con resorte tienen balasto lateral",
          all(s.u1 > 0 and s.u2 > 0 for s in m.springs), True)

    # con la opcion apagada vuelve el resorte de punta y desaparecen las restricciones
    p.piles.tip_restraint = False
    m2 = build(p).model
    check("sin la opcion no hay restricciones", len(m2.restraints), 0, 0.0)
    check("y reaparece el resorte de punta",
          sum(1 for s in m2.springs if s.u3 > 0), n_pilas, 0.0)


# --------------------------------------------------------------------------
def main() -> int:
    test_torsion()
    test_flexion()
    p, res, checks = test_franjas()
    test_datos_para_la_interfaz(checks)
    test_torsion_en_viga(checks)
    test_casos_analisis()
    test_apoyos_de_pila()
    test_reporte(p, res, checks)
    print("\n%s" % ("TODO OK" if not FAILS else "FALLAS (%d): %s" % (len(FAILS), ", ".join(FAILS))))
    return 1 if FAILS else 0


if __name__ == "__main__":
    raise SystemExit(main())
