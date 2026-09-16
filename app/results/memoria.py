"""Memoria de calculo en Word (.docx).

El objetivo es que salga firmable: no un volcado de resultados, sino el
desarrollo con las formulas y los numeros sustituidos, de modo que un revisor
pueda seguir el calculo sin abrir el modelo.

Se apoya en `python-docx`, que ya viene en requirements.
"""
from __future__ import annotations

import datetime as _dt

import io as _io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..core.params import WallProject
from ..exporters import raster
from . import detailing as DET
from . import rebar as R

AZUL = RGBColor(0x1F, 0x38, 0x64)
ROJO = RGBColor(0xC0, 0x30, 0x30)
GRIS = RGBColor(0x55, 0x55, 0x55)


# --------------------------------------------------------------------------
# Utilidades de formato
# --------------------------------------------------------------------------
def _n(x, d: int = 2, guion: str = "—") -> str:
    if x is None:
        return guion
    if isinstance(x, bool):
        return "SI" if x else "NO"
    if isinstance(x, (int, float)):
        return ("%%.%df" % d) % x
    return str(x)


def _campo(par, instruccion: str) -> None:
    """Inserta un campo de Word (indice, numero de pagina...) en el parrafo.

    Word rellena estos campos al abrir o al actualizar; no se pueden calcular
    aqui porque dependen de la paginacion real.
    """
    ini = OxmlElement("w:fldChar")
    ini.set(qn("w:fldCharType"), "begin")
    txt = OxmlElement("w:instrText")
    txt.set(qn("xml:space"), "preserve")
    txt.text = instruccion
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    r = par.add_run()._r
    for e in (ini, txt, sep, fin):
        r.append(e)


class _Doc:
    """Envoltorio fino sobre python-docx con los estilos de la memoria."""

    def __init__(self, titulo: str) -> None:
        self.d = Document()
        self._pagina()
        self._estilos()
        self.n_cap = 0
        self.n_sec = 0
        self.n_fig = 0
        self.n_tab = 0
        self.titulo = titulo
        self._encabezado()

    def _pagina(self) -> None:
        s = self.d.sections[0]
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)

    def _estilos(self) -> None:
        n = self.d.styles["Normal"]
        n.font.name = "Calibri"
        n.font.size = Pt(10.5)
        n.paragraph_format.space_after = Pt(5)

    def _encabezado(self) -> None:
        sec = self.d.sections[0]
        enc = sec.header.paragraphs[0]
        enc.text = self.titulo
        enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in enc.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = GRIS

        pie = sec.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pie.add_run("Pagina ")
        r.font.size = Pt(8)
        r.font.color.rgb = GRIS
        _campo(pie, "PAGE")
        r = pie.add_run(" de ")
        r.font.size = Pt(8)
        r.font.color.rgb = GRIS
        _campo(pie, "NUMPAGES")
        for r in pie.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = GRIS

    def indice(self) -> None:
        h = self.d.add_heading(level=1)
        r = h.add_run("Indice")
        r.font.color.rgb = AZUL
        r.font.size = Pt(15)
        par = self.d.add_paragraph()
        _campo(par, r'TOC \\o "1-2" \\h \\z \\u')
        self.p("Para ver el indice: clic derecho sobre el y «Actualizar campos». "
               "Word lo rellena con la paginacion real.", cursiva=True, color=GRIS)
        self.d.add_page_break()

    def figura(self, escena, pie: str, ancho_cm: float = 15.0,
               ancho_px: int = 1500) -> None:
        """Rasteriza una escena de dibujo y la incrusta con su pie."""
        if escena is None or not escena.prims:
            return
        try:
            png = raster.to_png(escena, ancho_px=ancho_px)
        except Exception:
            return          # una figura no puede tumbar la memoria
        self.n_fig += 1
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(_io.BytesIO(png), width=Cm(ancho_cm))
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run("Figura %d. %s" % (self.n_fig, pie))
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS

    # -- bloques -----------------------------------------------------------
    def portada(self, p: WallProject) -> None:
        for _ in range(4):
            self.d.add_paragraph()
        t = self.d.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run("MEMORIA DE CÁLCULO")
        r.bold = True
        r.font.size = Pt(26)
        r.font.color.rgb = AZUL

        s = self.d.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = s.add_run("Muro de contención sobre pilas con viga cabezal")
        r.font.size = Pt(14)
        r.font.color.rgb = GRIS

        self.d.add_paragraph()
        n = self.d.add_paragraph()
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = n.add_run(p.info.name)
        r.bold = True
        r.font.size = Pt(16)

        for _ in range(6):
            self.d.add_paragraph()

        tb = self.d.add_table(rows=0, cols=2)
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in (("Ingeniero", p.info.engineer or "—"),
                     ("Matrícula profesional", p.info.license or "—"),
                     ("Cliente", p.info.client or "—"),
                     ("Norma de diseño", "ACI 318-19"),
                     ("Modelo estructural", "SAP2000 v27"),
                     ("Fecha", _dt.date.today().strftime("%d/%m/%Y"))):
            row = tb.add_row().cells
            row[0].paragraphs[0].add_run(k + ":").bold = True
            row[1].text = str(v)
        if p.info.notes:
            self.d.add_paragraph()
            self.p(p.info.notes, cursiva=True)
        self.d.add_page_break()

    def cap(self, texto: str) -> None:
        self.n_cap += 1
        self.n_sec = 0
        h = self.d.add_heading(level=1)
        r = h.add_run("%d. %s" % (self.n_cap, texto))
        r.font.color.rgb = AZUL
        r.font.size = Pt(15)

    def sec(self, texto: str) -> None:
        self.n_sec += 1
        h = self.d.add_heading(level=2)
        r = h.add_run("%d.%d  %s" % (self.n_cap, self.n_sec, texto))
        r.font.color.rgb = AZUL
        r.font.size = Pt(12)

    def p(self, texto: str, cursiva: bool = False, color=None) -> None:
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = par.add_run(texto)
        r.italic = cursiva
        if color is not None:
            r.font.color.rgb = color

    def formula(self, expr: str, resultado: str = "") -> None:
        """Una linea de desarrollo, en monoespaciada y sangrada."""
        par = self.d.add_paragraph()
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(expr)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        if resultado:
            r2 = par.add_run("   =  " + resultado)
            r2.font.name = "Consolas"
            r2.font.size = Pt(9.5)
            r2.bold = True

    def kv(self, filas: list[tuple[str, object]], anchos=(9.0, 7.0)) -> None:
        tb = self.d.add_table(rows=0, cols=2)
        tb.style = "Light List Accent 1"
        for k, v in filas:
            c = tb.add_row().cells
            c[0].text = str(k)
            c[1].text = str(v)
        for row in tb.rows:
            row.cells[0].width = Cm(anchos[0])
            row.cells[1].width = Cm(anchos[1])
        self.d.add_paragraph()

    def tabla(self, cabecera: list[str], filas: list[list], nota: str = "",
              ancho_pt: float = 8.0) -> None:
        tb = self.d.add_table(rows=1, cols=len(cabecera))
        tb.style = "Light Grid Accent 1"
        for i, h in enumerate(cabecera):
            cel = tb.rows[0].cells[i]
            cel.text = ""
            r = cel.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(ancho_pt)
        for f in filas:
            c = tb.add_row().cells
            for i, v in enumerate(f):
                c[i].text = ""
                r = c[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(ancho_pt)
        if nota:
            self.p(nota, cursiva=True)
        self.d.add_paragraph()

    def aviso(self, texto: str) -> None:
        par = self.d.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.5)
        r = par.add_run("⚠  " + texto)
        r.font.color.rgb = ROJO
        r.font.size = Pt(10)

    def save(self, path: str) -> str:
        self.d.save(path)
        return path


# --------------------------------------------------------------------------
# Capitulos
# --------------------------------------------------------------------------
def _cap_datos(doc: _Doc, p: WallProject, report: dict, sch: dict | None) -> None:
    doc.cap("Datos de partida")
    if sch:
        doc.figura(DET.planta(p, sch),
                   "Planta de replanteo: ejes, pilas y viga cabezal.", 15.5, 1800)
    g = report.get("geometria", {})
    doc.sec("Geometría")
    doc.kv([
        ("Altura de la pantalla", "%s m" % _n(g.get("altura_muro"))),
        ("Longitud del muro", "%s m" % _n(g.get("longitud_muro"))),
        ("Espesor de la pantalla", "%s m" % _n(g.get("espesor_muro"))),
        ("Número de pilas", g.get("n_pilas")),
        ("Diámetro de pila", "%s m" % _n(g.get("diametro_pila"))),
        ("Separación entre pilas", "%s m" % _n(g.get("separacion_pilas"))),
        ("Longitud de pila", "%s m" % _n(g.get("longitud_pila"))),
        ("Viga cabezal", "%s × %s m" % (_n(p.cap_beam.width), _n(p.cap_beam.depth))),
    ])

    doc.sec("Materiales")
    c = p.materials.concrete
    doc.kv([
        ("Concreto f'c", "%s MPa" % _n(c.fc / 1000.0, 1)),
        ("Módulo E", "%s MPa" % _n(c.E_calc / 1000.0, 0)),
        ("Peso específico del concreto", "%s kN/m³" % _n(c.gamma, 1)),
        ("Acero de refuerzo fy", "%s MPa" % _n(p.materials.rebar.fy / 1000.0, 0)),
    ])
    doc.p("El módulo se obtiene de la expresión de ACI 318 19.2.2.1 para concreto "
          "de peso normal, salvo que se haya impuesto un valor.")
    doc.formula("E = 4700 · √f'c  [MPa]",
                "%s MPa" % _n(c.E_calc / 1000.0, 0))

    doc.sec("Perfil de suelo y balasto")
    doc.tabla(
        ["Estrato", "Z sup [m]", "Z inf [m]", "ks_h [kN/m³]", "ks_v [kN/m³]",
         "k lateral nodal [kN/m]"],
        [[l["estrato"], _n(l["z_top"]), _n(l["z_bot"]), _n(l["ks_h"], 0),
          _n(l["ks_v"], 0), _n(l["k_lat_nodal"], 0)]
         for l in report.get("balasto", [])],
        nota="El resorte lateral de cada nudo es ks_h · D · Δz, con Δz la longitud "
             "tributaria del nudo sobre la pila. Los nudos de arranque y punta toman "
             "media longitud tributaria.")
    if p.soil.water_table_z is not None:
        doc.p("Nivel freático a la cota Z = %s m; por debajo se usa el peso específico "
              "sumergido y se añade el empuje hidrostático." % _n(p.soil.water_table_z))


def _cap_empujes(doc: _Doc, p: WallProject, report: dict) -> None:
    doc.cap("Empujes de tierras")
    e = report.get("empujes", {})
    pr = report.get("presiones_kPa", {})

    doc.sec("Coeficiente de empuje")
    doc.p("Método: %s, condición %s." % (p.earth.method, p.earth.condition))
    if p.earth.method == "rankine" and p.earth.condition == "activo":
        doc.formula("Ka = tan²(45° − φ/2)   con φ = %s°" % _n(p.earth.phi, 1),
                    _n(e.get("K"), 4))
    elif p.earth.condition == "reposo":
        doc.formula("K0 = 1 − sen φ   con φ = %s°" % _n(p.earth.phi, 1),
                    _n(e.get("K"), 4))
    else:
        doc.formula("K (%s, %s)" % (p.earth.method, p.earth.condition),
                    _n(e.get("K"), 4))
    doc.formula("K · γ = %s · %s" % (_n(e.get("K"), 4), _n(p.earth.gamma, 1)),
                "%s kN/m³" % _n(e.get("Ka_gamma"), 3))

    if p.seismic.enabled:
        doc.sec("Incremento sísmico")
        doc.p("Se aplica el incremento dinámico de Mononobe-Okabe: ΔKae = Kae − Ka, "
              "repartido en altura según la distribución «%s»."
              % p.seismic.distribution)
        doc.formula("kh = %s   kv = %s" % (_n(p.seismic.kh, 3), _n(p.seismic.kv, 3)))
        doc.formula("ΔKae", _n(e.get("dK"), 4))
        doc.formula("ΔKae · γ", "%s kN/m³" % _n(e.get("dKae_gamma"), 3))

    doc.figura(DET.diagrama_empujes(p, report),
               "Diagrama de presiones sobre la pantalla, por componentes.", 13.0)
    doc.sec("Presiones y resultantes")
    doc.kv([
        ("Empuje del suelo en la base", "%s kPa" % _n(pr.get("suelo_base"))),
        ("Incremento sísmico en la corona", "%s kPa" % _n(pr.get("sismo_corona"))),
        ("Sobrecarga", "%s kPa" % _n(pr.get("sobrecarga"))),
    ])
    doc.tabla(
        ["Acción", "Resultante [kN/m]", "Brazo [m]"],
        [["Empuje del suelo", _n(e.get("P_suelo"), 1), _n(e.get("y_suelo"))],
         ["Sobrecarga", _n(e.get("P_sobrecarga"), 1), _n(e.get("y_sobrecarga"))],
         ["Incremento sísmico", _n(e.get("P_sismo"), 1), _n(e.get("y_sismo"))],
         ["Agua", _n(e.get("P_agua"), 1), "—"],
         ["TOTAL", _n(e.get("P_total"), 1), "—"]],
        nota="Resultantes por metro lineal de muro. Los brazos se miden desde la base "
             "de la pantalla.")


def _cap_modelo(doc: _Doc, p: WallProject, report: dict, checks: dict | None) -> None:
    doc.cap("Modelo estructural")
    mo = report.get("modelo", {})
    doc.p("El modelo se genera en SAP2000. La pantalla se discretiza con elementos "
          "shell, las pilas y la viga cabezal con elementos frame, y el suelo con "
          "resortes de balasto por estrato en los nudos de las pilas.")
    doc.kv([
        ("Nudos", mo.get("joints")),
        ("Elementos frame", mo.get("frames")),
        ("Elementos shell", mo.get("areas")),
        ("Resortes de balasto", mo.get("springs")),
        ("Patrones de carga", ", ".join(mo.get("load_patterns", []))),
        ("Combinaciones", str(len(mo.get("combos", [])))),
        ("Divisiones de shell por vano", p.mesh.wall_div_per_bay),
        ("Altura objetivo de shell", "%s m" % _n(p.mesh.wall_target_dz)),
    ])

    doc.sec("Combinaciones de carga")
    doc.tabla(["Combinación", "Factores", "Uso"],
              [[c.name, "  +  ".join("%s·%s" % (_n(f.factor, 2), f.pattern)
                                     for f in c.factors),
                "Diseño" if c.design == "Strength" else "Servicio"]
               for c in p.combos])

    if checks and checks.get("modal"):
        md = checks["modal"]
        doc.sec("Análisis modal")
        doc.kv([
            ("Número de modos", md["n_modos"]),
            ("Periodo fundamental T₁", "%s s" % _n(md["T1"], 4)),
            ("Modo dominante en X", "%s (T = %s s)"
             % (md["modo_dominante_X"], _n(md["T_dominante_X"], 4))),
            ("Masa acumulada en X", _n(md["masa_acumulada_X"], 3)),
            ("Masa acumulada en Y", _n(md["masa_acumulada_Y"], 3)),
        ])
        if md["masa_acumulada_X"] < 0.9 and md["masa_acumulada_Y"] < 0.9:
            doc.aviso("La masa participante acumulada no alcanza 0.90 en ninguna "
                      "dirección: conviene ampliar el número de modos.")


def _cap_verificaciones(doc: _Doc, p: WallProject, checks: dict) -> None:
    doc.cap("Verificación de los elementos")
    res = checks.get("resumen", {})

    doc.sec("Pilas")
    fichas = checks.get("pilas_resumen") or []
    if fichas:
        doc.p("Las pilas no son intercambiables: cada una tiene su longitud, su "
              "cota de cabeza y su demanda. La tabla siguiente resume lo que "
              "gobierna el armado de cada una.")
        doc.tabla(
            ["Pila", "D [m]", "Cabeza", "Punta", "L [m]", "N [kN]", "M [kN·m]",
             "V [kN]", "D/C P-M", "Cumple"],
            [[str(f["pila"]), _n(f["diametro"]), _n(f["z_cabeza"]), _n(f["z_punta"]),
              _n(f["longitud"]), _n(f["Pu_max"], 0), _n(f["Mu_max"], 0),
              _n(f["Vu_max"], 0), _n(f["ratio_PM"], 2),
              "Si" if f["ok"] else "NO"] for f in fichas])

    doc.figura(DET.diagrama_interaccion(checks),
               "Interaccion P-M de la zona mas solicitada, con la demanda de "
               "cada elemento de pila.", 12.5)
    doc.p("Cada pila se comprueba con el diagrama de interacción P-M de la sección "
          "circular, obtenido por compatibilidad de deformaciones con φ variable "
          "según la deformación neta de tracción (ACI 318 21.2.2) y tope 0.80·φ·Po. "
          "El cortante sigue 22.5 con bw = D y d = 0.8·D.")
    peores = sorted(checks.get("pilas", []),
                    key=lambda r: max(r["ratio_PM"], r["ratio_V"]), reverse=True)[:10]
    doc.tabla(
        ["Frame", "Z [m]", "Sección", "Pu [kN]", "Mu [kN·m]", "φMn [kN·m]",
         "D/C P-M", "Vu [kN]", "D/C V", "ρ", "OK"],
        [[r["elemento"], _n(r["z"], 1), r["seccion"], _n(r["Pu"], 0), _n(r["Mu"], 0),
          _n(r["phiMn"], 0), _n(r["ratio_PM"], 3), _n(r["Vu"], 0),
          _n(r["ratio_V"], 3), _n(r["rho"], 4), "SÍ" if r["ok"] else "NO"]
         for r in peores],
        nota="Diez elementos más solicitados. D/C máximo en pilas: %s."
             % _n(res.get("ratio_max_pilas"), 3), ancho_pt=7.5)

    doc.sec("Viga cabezal")
    doc.p("Flexión rectangular con φ iterado, cortante con estribos según 22.5 y "
          "torsión según 22.7. La torsión de esta viga es de equilibrio —nace del "
          "voladizo de la pantalla y no tiene camino alterno—, por lo que no se "
          "aplica la redistribución de 22.7.5.")
    vig = sorted(checks.get("viga_cabezal", []),
                 key=lambda r: max(r["ratio_M"], r["ratio_V"]), reverse=True)[:8]
    doc.tabla(
        ["Frame", "Mu [kN·m]", "As req [mm²]", "D/C M", "Vu [kN]", "D/C V",
         "Tu [kN·m]", "D/C V+T", "OK"],
        [[r["elemento"], _n(r["Mu"], 0), _n(r["As_req"], 0), _n(r["ratio_M"], 3),
          _n(r["Vu"], 0), _n(r["ratio_V"], 3), _n(r.get("Tu"), 0),
          _n(r.get("ratio_VT"), 3), "SÍ" if r["ok"] else "NO"]
         for r in vig], ancho_pt=7.5)

    tor = checks.get("torsion_viga")
    if tor:
        doc.p("Desarrollo del diseño por torsión en la sección crítica "
              "(elemento %s):" % tor["elemento_critico"])
        doc.formula("Tu", "%s kN·m" % _n(tor["Tu"], 1))
        doc.formula("T umbral = φ·0.083·λ·√f'c·Acp²/pcp   (22.7.4)",
                    "%s kN·m" % _n(tor["T_umbral"], 1))
        doc.formula("Tcr = 0.33·λ·√f'c·Acp²/pcp   (22.7.5)",
                    "%s kN·m" % _n(tor["T_cr"], 1))
        doc.p("Como Tu supera el umbral, la torsión no es despreciable y se diseña:")
        doc.formula("Aoh = %s mm²    ph = %s mm    Ao = 0.85·Aoh = %s mm²"
                    % (_n(tor.get("Aoh"), 0), _n(tor.get("ph"), 0),
                       _n(tor.get("Ao"), 0)))
        doc.formula("At/s = Tu / (φ·2·Ao·fyt·cotθ)   con θ = 45°",
                    "%s mm²/m por rama" % _n(tor["At_s"], 0))
        doc.formula("(Av + 2·At)/s  requerido",
                    "%s mm²/m" % _n(tor["Avt_s"], 0))
        doc.formula("(Av + 2·At)/s  mínimo   (9.6.4.2)",
                    "%s mm²/m" % _n(tor["Avt_s_min"], 0))
        doc.formula("Al = (At/s)·ph·(fyt/fy)·cot²θ   (22.7.6.1b)",
                    "%s mm²" % _n(tor["Al"], 0))
        doc.formula("s ≤ mín(ph/8, 300 mm)   (9.7.6.3.3)",
                    "%s mm" % _n(tor["s_max"], 0))
        doc.formula("D/C del límite de sección V+T   (22.7.7.1)",
                    _n(tor["ratio_VT"], 3))
        doc.p("Resultado: estribos cerrados %s @ %s mm más %d barras %s repartidas en "
              "el perímetro a no más de 300 mm, adicionales al refuerzo por flexión."
              % (tor["estribo"], _n(tor["s_estribo"], 0), tor["n_barras_torsion"],
                 tor["barra_torsion"]))
        if not tor["seccion_ok"]:
            doc.aviso("La sección no admite la combinación de cortante y torsión "
                      "(ACI 318 22.7.7.1): hay que ampliar el ancho o el canto.")

    cmp_t = checks.get("comparacion_torsion")
    if cmp_t:
        doc.p("Contraste con el diseño de concreto de SAP2000 (%s):"
              % cmp_t.get("codigo_sap"))
        doc.tabla(["Concepto", "Cálculo propio", "SAP2000", "Diferencia"],
                  [["Al por torsión [mm²]", _n(cmp_t["Al_propio"], 0),
                    _n(cmp_t["Al_sap"], 0), "%s %%" % _n(cmp_t["dif_Al_pct"], 1)],
                   ["At/s por torsión [mm²/m]", _n(cmp_t["At_s_propio"], 0),
                    _n(cmp_t["At_s_sap"], 0), "%s %%" % _n(cmp_t["dif_At_s_pct"], 1)]],
                  nota=cmp_t.get("nota", ""))

    doc.sec("Pantalla")
    muro = checks.get("muro") or {}
    fr = (muro.get("franjas") or {})
    doc.p("Los momentos de shell se promedian sobre los nudos de cada elemento: los "
          "valores nodales sobre un apoyo puntual son singularidades numéricas que "
          "crecen al refinar la malla y no son solicitación de diseño. En el modelo "
          "M22 pasa de %s kN·m/m en el pico nodal a %s kN·m/m promediado."
          % (_n(muro.get("pico_M22"), 0), _n(muro.get("M22_max"), 0)))
    if fr:
        doc.p("El dimensionado se hace por franjas de diseño: la altura se parte en %d "
              "bandas y cada una se arma con su propia solicitación, promediada sobre "
              "sus elementos. Además se aplica ACI 318 9.4.2.1, que permite diseñar "
              "con el momento en la cara del apoyo: en vertical la cara de la viga "
              "cabezal (Z = %s m) y en horizontal la de la pila (± %s m del eje)."
              % (fr["n_franjas"], _n(fr["z_cara_viga"]), _n(fr["radio_pila"])))
        doc.tabla(
            ["Franja", "Z ini", "Z fin", "M22 dis.", "As vert", "M11 cara",
             "As h apoyo", "As h vano", "D/C V", "OK"],
            [[f["franja"], _n(f["z_ini"]), _n(f["z_fin"]), _n(f["M22_diseno"], 0),
              _n(f["vertical"]["As"], 0), _n(f["M11_cara"], 0),
              _n(f["horizontal_apoyo"]["As"], 0), _n(f["horizontal_vano"]["As"], 0),
              _n(f["cortante"]["ratio"], 3),
              "SÍ" if (f["vertical"]["ok"] and f["cortante"]["ok"]) else "NO"]
             for f in fr["franjas"]],
            nota="Momentos en kN·m/m y áreas en mm²/m.", ancho_pt=7.5)
        if fr.get("reduccion_cara_horizontal_pct"):
            doc.p("Tomar el momento en la cara de la pila en lugar de en su eje reduce "
                  "la solicitación horizontal un %s %%."
                  % _n(fr["reduccion_cara_horizontal_pct"], 1))
    doc.p("El D/C a flexión de la pantalla vale 1.00 por construcción, porque ahí se "
          "calcula el acero necesario en vez de comprobar uno dado. El único D/C con "
          "significado es el de cortante, que no lleva refuerzo transversal: %s."
          % _n(res.get("ratio_cortante_muro"), 3))

    if res.get("elementos_no_conformes"):
        doc.aviso("%d elemento(s) no cumplen la verificación. Revisar antes de "
                  "construir." % res["elementos_no_conformes"])


def _cap_despiece(doc: _Doc, p: WallProject, sch: dict) -> None:
    doc.cap("Despiece de refuerzo")
    doc.p("Las longitudes de desarrollo, ganchos y traslapos siguen el capítulo 25 de "
          "ACI 318-19, con barra comercial de %s m y solape clase B. Las dimensiones "
          "de doblado se miden al eje de la barra. La lámina de despiece se entrega "
          "aparte en DXF, SVG y PDF."
          % _n(R.LARGO_COMERCIAL, 0))

    doc.figura(DET.corte_conjunto(p, sch),
               "Corte transversal con el refuerzo dispuesto.", 9.0, 1100)
    doc.figura(DET.detalle_arranque(p, sch, {}),
               "Detalle del arranque de la pantalla sobre la viga cabezal.", 13.0)
    doc.sec("Planilla de aceros")
    doc.tabla(
        ["Marca", "Elemento", "Posición", "Ø", "Forma", "Largo [m]", "Cant.",
         "Peso [kg]"],
        [[m["marca"], m["elemento"], m["posicion"], m["diametro"], m["forma"],
          _n(m["largo_m"], 2), m["cantidad"], _n(m["peso_total_kg"], 1)]
         for m in sch["marcas"]], ancho_pt=7.0)

    doc.sec("Resumen de materiales")
    doc.tabla(["Ø", "db [mm]", "Longitud total [m]", "Peso [kg]"],
              [[d["diametro"], _n(d["db_mm"], 1), _n(d["largo_m"], 0),
                _n(d["peso_kg"], 0)] for d in sch["por_diametro"]])
    doc.kv([
        ("Acero total", "%s kg" % _n(sch["peso_total_kg"], 0)),
        ("Concreto total", "%s m³" % _n(sch["volumen_concreto_m3"], 1)),
        ("Cuantía", "%s kg/m³" % _n(sch["cuantia_kg_m3"], 1)),
    ] + [("Acero en %s" % k, "%s kg" % _n(v, 0))
         for k, v in sch["por_elemento"].items()])


def _cap_conclusiones(doc: _Doc, checks: dict, warnings: list[str] | None) -> None:
    doc.cap("Conclusiones")
    res = checks.get("resumen", {}) if checks else {}
    malos = res.get("elementos_no_conformes", 0)
    if malos:
        doc.p("El modelo presenta %d elemento(s) que no cumplen la verificación. "
              "No procede construir hasta resolverlos." % malos)
    else:
        doc.p("Todos los elementos verificados cumplen las comprobaciones de ACI 318-19 "
              "para las combinaciones consideradas.")
    doc.kv([
        ("D/C máximo en pilas", _n(res.get("ratio_max_pilas"), 3)),
        ("D/C máximo en viga cabezal", _n(res.get("ratio_max_viga"), 3)),
        ("D/C de cortante en la pantalla", _n(res.get("ratio_cortante_muro"), 3)),
        ("Torsión significativa en la viga",
         "SÍ" if res.get("torsion_viga_significativa") else "NO"),
    ])
    if warnings:
        doc.sec("Advertencias del modelo")
        for w in warnings:
            doc.aviso(w)
    if checks.get("lecturas_fallidas"):
        doc.sec("Lecturas de resultados con problemas")
        for w in checks["lecturas_fallidas"]:
            doc.aviso(w)

    doc.d.add_paragraph()
    doc.p("Esta memoria se ha generado automáticamente a partir del modelo de SAP2000. "
          "Debe ser revisada y firmada por el ingeniero responsable antes de su uso.",
          cursiva=True, color=GRIS)


def _cap_anexos(doc: _Doc, p: WallProject, checks: dict) -> None:
    """Listados completos: lo que no cabe en el cuerpo pero hay que poder mirar."""
    env = checks.get("envolventes") or {}
    if not env:
        return
    doc.d.add_page_break()
    doc.cap("Anexo. Envolventes de solicitaciones")
    doc.p("Envolvente de cada elemento sobre las combinaciones de diseño, con la "
          "combinación que gobierna cada componente. Es el dato crudo del que "
          "salen las verificaciones del capítulo anterior.")

    viga = env.get("viga_cabezal") or []
    if viga:
        doc.sec("Viga cabezal")
        doc.tabla(
            ["Frame", "Y [m]", "P [kN]", "V2 [kN]", "V3 [kN]", "T [kN·m]",
             "M2 [kN·m]", "M3 [kN·m]", "Gobierna M"],
            [[r["elemento"], _n(r["y"]), _n(r["P"]["abs"], 0), _n(r["V2"]["abs"], 0),
              _n(r["V3"]["abs"], 0), _n(r["T"]["abs"], 0), _n(r["M2"]["abs"], 0),
              _n(r["M3"]["abs"], 0),
              r["M2"]["caso"] if r["M2"]["abs"] >= r["M3"]["abs"] else r["M3"]["caso"]]
             for r in sorted(viga, key=lambda x: x.get("y") or 0.0)], ancho_pt=7.0)

    pilas = env.get("pilas") or []
    if pilas:
        doc.sec("Pilas")
        peores = sorted(pilas, key=lambda r: max(r["M2"]["abs"], r["M3"]["abs"]),
                        reverse=True)[:20]
        doc.tabla(
            ["Frame", "Z [m]", "P [kN]", "V2 [kN]", "T [kN·m]", "M2 [kN·m]",
             "M3 [kN·m]", "Gobierna M"],
            [[r["elemento"], _n(r["z"], 1), _n(r["P"]["abs"], 0), _n(r["V2"]["abs"], 0),
              _n(r["T"]["abs"], 0), _n(r["M2"]["abs"], 0), _n(r["M3"]["abs"], 0),
              r["M2"]["caso"] if r["M2"]["abs"] >= r["M3"]["abs"] else r["M3"]["caso"]]
             for r in peores],
            nota="Veinte elementos con mayor momento.", ancho_pt=7.0)

    reac = checks.get("reaccion_vertical") or {}
    if reac:
        doc.sec("Reacción vertical por combinación")
        doc.tabla(["Combinación", "Suma Fz [kN]"],
                  [[k, _n(v, 1)] for k, v in reac.items()],
                  nota="Con el modelo en equilibrio, la reacción de la combinación "
                       "de servicio iguala el peso propio.")


# --------------------------------------------------------------------------
# Entrada publica
# --------------------------------------------------------------------------
def build_memoria(p: WallProject, report: dict, checks: dict | None = None,
                  schedule: dict | None = None,
                  warnings: list[str] | None = None) -> Document:
    doc = _Doc(p.info.name)
    doc.portada(p)
    doc.indice()

    doc.cap("Objeto y alcance")
    doc.p("Esta memoria recoge el cálculo de un muro de contención resuelto con una "
          "pantalla de concreto reforzado apoyada sobre una viga cabezal y pilas "
          "perforadas. Comprende los datos de partida, los empujes de tierras, el "
          "modelo estructural, la verificación de cada elemento según ACI 318-19 y el "
          "despiece del refuerzo.")
    doc.p("El análisis se realiza con SAP2000. El suelo se representa con resortes de "
          "balasto por estrato en las pilas; el empuje de tierras se aplica como "
          "presión superficial sobre la pantalla, modulada nodo a nodo por un patrón "
          "de junta, de modo que el diagrama real de presiones se reproduce sin "
          "escalonamientos.")

    _cap_datos(doc, p, report, schedule)
    _cap_empujes(doc, p, report)
    _cap_modelo(doc, p, report, checks)
    if checks:
        _cap_verificaciones(doc, p, checks)
    if schedule:
        _cap_despiece(doc, p, schedule)
    if checks:
        _cap_conclusiones(doc, checks, warnings)
        _cap_anexos(doc, p, checks)
    return doc.d


def save_memoria(p: WallProject, report: dict, path: str, checks: dict | None = None,
                 schedule: dict | None = None,
                 warnings: list[str] | None = None) -> str:
    build_memoria(p, report, checks, schedule, warnings).save(path)
    return path
