"""Reporte a Excel de datos de entrada, empujes y verificaciones."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.params import WallProject

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")


def _headers(ws, row: int, labels: list[str]) -> None:
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _autofit(ws, max_width: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), max_width)


def _kv_sheet(ws, title: str, pairs: list[tuple[str, object]], start: int = 1) -> int:
    ws.cell(row=start, column=1, value=title).font = TITLE_FONT
    r = start + 1
    for k, v in pairs:
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
        r += 1
    return r + 1


def build_report(project: WallProject, report: dict, checks: dict | None = None,
                 warnings: list[str] | None = None,
                 schedule: dict | None = None) -> Workbook:
    wb = Workbook()

    # ---- Datos -----------------------------------------------------------
    ws = wb.active
    ws.title = "Datos"
    g = report.get("geometria", {})
    r = _kv_sheet(ws, "Proyecto", [
        ("Nombre", project.info.name),
        ("Ingeniero", project.info.engineer),
        ("Matricula profesional", project.info.license),
        ("Cliente", project.info.client),
    ])
    r = _kv_sheet(ws, "Geometria", [
        ("Altura del muro [m]", g.get("altura_muro")),
        ("Longitud del muro [m]", g.get("longitud_muro")),
        ("Espesor del muro [m]", g.get("espesor_muro")),
        ("Numero de pilas", g.get("n_pilas")),
        ("Separacion de pilas [m]", g.get("separacion_pilas")),
        ("Diametro de pila [m]", g.get("diametro_pila")),
        ("Longitud de pila [m]", g.get("longitud_pila")),
        ("Viga cabezal [m]", "%.2f x %.2f" % (project.cap_beam.width, project.cap_beam.depth)),
    ], r)
    c = project.materials.concrete
    _kv_sheet(ws, "Materiales", [
        ("f'c [kPa]", c.fc),
        ("E [kPa]", round(c.E_calc, 0)),
        ("Peso especifico [kN/m3]", c.gamma),
        ("fy [kPa]", project.materials.rebar.fy),
    ], r)
    _autofit(ws)

    # ---- Empujes ---------------------------------------------------------
    ws = wb.create_sheet("Empujes")
    e = report.get("empujes", {})
    p = report.get("presiones_kPa", {})
    r = _kv_sheet(ws, "Coeficientes y resultantes (por metro de muro)", [
        ("Metodo", project.earth.method),
        ("Condicion", project.earth.condition),
        ("K", e.get("K")),
        ("gamma relleno [kN/m3]", project.earth.gamma),
        ("K * gamma [kN/m3]", e.get("Ka_gamma")),
        ("Sobrecarga [kPa]", project.earth.surcharge),
        ("Nivel freatico [m]", project.soil.water_table_z),
        ("", ""),
        ("Sismo habilitado", project.seismic.enabled),
        ("kh", project.seismic.kh),
        ("kv", project.seismic.kv),
        ("dKae", e.get("dK")),
        ("dKae * gamma [kN/m3]", e.get("dKae_gamma")),
        ("Distribucion sismica", project.seismic.distribution),
    ])
    r = _kv_sheet(ws, "Resultantes [kN/m]", [
        ("P suelo", e.get("P_suelo")), ("brazo suelo [m]", e.get("y_suelo")),
        ("P sobrecarga", e.get("P_sobrecarga")), ("brazo sobrecarga [m]", e.get("y_sobrecarga")),
        ("P sismo", e.get("P_sismo")), ("brazo sismo [m]", e.get("y_sismo")),
        ("P agua", e.get("P_agua")),
        ("P total", e.get("P_total")),
    ], r)
    _kv_sheet(ws, "Presiones caracteristicas [kPa]", [
        ("Empuje en la base", p.get("suelo_base")),
        ("Incremento sismico en corona", p.get("sismo_corona")),
        ("Sobrecarga", p.get("sobrecarga")),
    ], r)
    _autofit(ws)

    # ---- Balasto ---------------------------------------------------------
    ws = wb.create_sheet("Balasto")
    ws.cell(row=1, column=1, value="Perfil de suelo y resortes").font = TITLE_FONT
    _headers(ws, 2, ["Estrato", "Z sup [m]", "Z inf [m]", "ks_h [kN/m3]", "ks_v [kN/m3]",
                     "k lateral nodal [kN/m]"])
    for i, lay in enumerate(report.get("balasto", []), start=3):
        ws.cell(row=i, column=1, value=lay["estrato"])
        ws.cell(row=i, column=2, value=lay["z_top"])
        ws.cell(row=i, column=3, value=lay["z_bot"])
        ws.cell(row=i, column=4, value=lay["ks_h"])
        ws.cell(row=i, column=5, value=lay["ks_v"])
        ws.cell(row=i, column=6, value=lay["k_lat_nodal"])
    _autofit(ws)

    # ---- Modelo ----------------------------------------------------------
    ws = wb.create_sheet("Modelo")
    mo = report.get("modelo", {})
    r = _kv_sheet(ws, "Topologia", [
        ("Nudos", mo.get("joints")), ("Frames", mo.get("frames")),
        ("Shells", mo.get("areas")), ("Resortes", mo.get("springs")),
        ("Patrones de carga", ", ".join(mo.get("load_patterns", []))),
        ("Combinaciones", ", ".join(mo.get("combos", []))),
        ("Grupos", ", ".join(mo.get("groups", []))),
    ])
    if warnings:
        ws.cell(row=r, column=1, value="Advertencias").font = TITLE_FONT
        for i, wmsg in enumerate(warnings, start=r + 1):
            ws.cell(row=i, column=1, value=wmsg)
    _autofit(ws)

    if not checks:
        return wb

    # ---- Resumen por pila ------------------------------------------------
    # Va antes que el detalle por elemento porque es la tabla que se mira: una
    # fila por pila, con su geometria y lo que gobierna su armado.
    fichas = checks.get("pilas_resumen") or []
    if fichas:
        ws = wb.create_sheet("PilasResumen")
        _headers(ws, 1, ["Pila", "Y [m]", "D [m]", "Cabeza Z [m]", "Punta Z [m]",
                         "Longitud [m]", "Pu [kN]", "Mu [kN-m]", "Z de Mu [m]",
                         "Vu [kN]", "D/C P-M", "D/C V", "rho", "Zonas", "OK"])
        for i, f in enumerate(fichas, start=2):
            vals = [f["pila"], f["y"], f["diametro"], f["z_cabeza"], f["z_punta"],
                    f["longitud"], f["Pu_max"], f["Mu_max"], f["z_Mu_max"],
                    f["Vu_max"], f["ratio_PM"], f["ratio_V"], f["rho"],
                    ", ".join(f["zonas"]), "SI" if f["ok"] else "NO"]
            for cix, v in enumerate(vals, start=1):
                ws.cell(row=i, column=cix, value=v)
            ws.cell(row=i, column=15).fill = OK_FILL if f["ok"] else BAD_FILL
        _autofit(ws)

    # ---- Pilas, elemento a elemento --------------------------------------
    ws = wb.create_sheet("Pilas")
    _headers(ws, 1, ["Frame", "Pila", "Y [m]", "Z [m]", "Seccion", "Pu [kN]", "Mu [kN-m]",
                     "phiMn [kN-m]", "D/C P-M", "Vu [kN]", "phiVn [kN]", "D/C V", "rho", "OK"])
    for i, rec in enumerate(checks.get("pilas", []), start=2):
        vals = [rec["elemento"], rec.get("pila"), rec["y"], rec["z"], rec["seccion"],
                rec["Pu"], rec["Mu"], rec["phiMn"], rec["ratio_PM"], rec["Vu"],
                rec["phiVn"], rec["ratio_V"], rec["rho"], "SI" if rec["ok"] else "NO"]
        for cix, v in enumerate(vals, start=1):
            ws.cell(row=i, column=cix, value=v)
        ws.cell(row=i, column=14).fill = OK_FILL if rec["ok"] else BAD_FILL
    _autofit(ws)

    # ---- Viga cabezal ----------------------------------------------------
    ws = wb.create_sheet("VigaCabezal")
    _headers(ws, 1, ["Frame", "Y [m]", "Mu [kN-m]", "As req [mm2]", "As min [mm2]",
                     "phiMn [kN-m]", "D/C M", "Vu [kN]", "Av/s req [mm2/m]", "phiVn [kN]",
                     "D/C V", "Tu [kN-m]", "T umbral [kN-m]", "Torsion",
                     "At/s [mm2/m]", "(Av+2At)/s [mm2/m]", "Al [mm2]",
                     "Estribo @ [mm]", "n barras torsion", "D/C V+T", "OK"])
    for i, rec in enumerate(checks.get("viga_cabezal", []), start=2):
        vals = [rec["elemento"], rec["y"], rec["Mu"], rec["As_req"], rec["As_min"],
                rec["phiMn"], rec["ratio_M"], rec["Vu"], rec["Av_s_req"], rec["phiVn"],
                rec["ratio_V"], rec.get("Tu"), rec.get("T_umbral"),
                "SIGNIFICATIVA" if rec.get("torsion_significativa") else "despreciable",
                rec.get("At_s"), rec.get("Avt_s"), rec.get("Al_adopt"),
                rec.get("s_estribo"), rec.get("n_barras_torsion"), rec.get("ratio_VT"),
                "SI" if rec["ok"] else "NO"]
        for cix, v in enumerate(vals, start=1):
            ws.cell(row=i, column=cix, value=v)
        if rec.get("torsion_significativa"):
            ws.cell(row=i, column=14).fill = BAD_FILL
        ws.cell(row=i, column=21).fill = OK_FILL if rec["ok"] else BAD_FILL
    _autofit(ws)

    # ---- Muro ------------------------------------------------------------
    ws = wb.create_sheet("Muro")
    mu = checks.get("muro", {})
    if mu:
        r = _kv_sheet(ws, "Solicitaciones criticas (promediadas por elemento)", [
            ("M11 max [kN-m/m]", mu["M11_max"]), ("Area critica M11", mu["area_M11"]),
            ("M22 max [kN-m/m]", mu["M22_max"]), ("Area critica M22", mu["area_M22"]),
            ("", ""),
            ("Pico nodal M11 [kN-m/m]", mu.get("pico_M11")),
            ("Pico nodal M22 [kN-m/m]", mu.get("pico_M22")),
            ("Pico nodal V [kN/m]", mu.get("pico_V")),
            ("Nota", mu.get("nota", "")),
        ])
        _headers(ws, r, ["Direccion", "As req [mm2/m]", "As min [mm2/m]", "As adoptado [mm2/m]",
                         "phiMn [kN-m/m]", "D/C", "OK"])
        for i, key in enumerate(("horizontal", "vertical"), start=r + 1):
            d = mu[key]
            for cix, v in enumerate([key, d["As_req"], d["As_min"], d["As"], d["phiMn"],
                                     d["ratio"], "SI" if d["ok"] else "NO"], start=1):
                ws.cell(row=i, column=cix, value=v)
            ws.cell(row=i, column=7).fill = OK_FILL if d["ok"] else BAD_FILL
        sv = mu["cortante"]
        _kv_sheet(ws, "Cortante", [
            ("Vu [kN/m]", sv["Vu"]), ("phiVc [kN/m]", sv["phiVc"]),
            ("D/C", sv["ratio"]), ("OK", "SI" if sv["ok"] else "NO"),
        ], r + 4)
    _autofit(ws)

    # ---- Torsion de la viga cabezal --------------------------------------
    tor = checks.get("torsion_viga") or {}
    cmp_tor = checks.get("comparacion_torsion") or {}
    if tor:
        ws = wb.create_sheet("Torsion")
        r = _kv_sheet(ws, "Diseno por torsion de la viga cabezal (ACI 318 22.7)", [
            ("Elemento critico", tor["elemento_critico"]),
            ("Y [m]", tor["y"]),
            ("Tu [kN-m]", tor["Tu"]),
            ("Umbral 22.7.4 [kN-m]", tor["T_umbral"]),
            ("Tcr agrietamiento [kN-m]", tor["T_cr"]),
            ("Elementos con torsion significativa", tor["n_elementos"]),
            ("", ""),
            ("D/C limite de seccion V+T (22.7.7.1)", tor["ratio_VT"]),
            ("Seccion admisible", "SI" if tor["seccion_ok"] else "NO"),
        ])
        r = _kv_sheet(ws, "Refuerzo requerido", [
            ("At/s por torsion, una rama [mm2/m]", tor["At_s"]),
            ("(Av+2At)/s total [mm2/m]", tor["Avt_s"]),
            ("Minimo 9.6.4.2 [mm2/m]", tor["Avt_s_min"]),
            ("Estribo cerrado", tor["estribo"]),
            ("Separacion requerida [mm]", tor["s_estribo"]),
            ("Separacion maxima 9.7.6.3.3 [mm]", tor["s_max"]),
            ("", ""),
            ("Al longitudinal por torsion [mm2]", tor["Al"]),
            ("Barra longitudinal", tor["barra_torsion"]),
            ("Numero de barras en el perimetro", tor["n_barras_torsion"]),
        ], r)
        ws.cell(row=r, column=1, value="Nota").font = TITLE_FONT
        ws.cell(row=r + 1, column=1, value=tor.get("nota", ""))
        ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True)
        r += 3
        if cmp_tor:
            r = _kv_sheet(ws, "Contraste con el diseno de SAP2000", [
                ("Codigo usado por SAP", cmp_tor.get("codigo_sap")),
                ("Al propio [mm2]", cmp_tor.get("Al_propio")),
                ("Al de SAP [mm2]", cmp_tor.get("Al_sap")),
                ("Diferencia Al [%]", cmp_tor.get("dif_Al_pct")),
                ("At/s propio [mm2/m]", cmp_tor.get("At_s_propio")),
                ("At/s de SAP [mm2/m]", cmp_tor.get("At_s_sap")),
                ("Diferencia At/s [%]", cmp_tor.get("dif_At_s_pct")),
            ], r)
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 20

    # ---- Franjas de diseno de la pantalla --------------------------------
    fr = (checks.get("muro") or {}).get("franjas") or {}
    if fr:
        ws = wb.create_sheet("MuroFranjas")
        r = _kv_sheet(ws, "Franjas de diseno de la pantalla", [
            ("Numero de franjas", fr["n_franjas"]),
            ("Filas de shells en altura", fr["n_filas"]),
            ("Cota de la cara de la viga cabezal [m]", fr["z_cara_viga"]),
            ("Radio de la pila [m]", fr["radio_pila"]),
            ("Reduccion vertical por cara [%]", fr.get("reduccion_cara_vertical_pct")),
            ("Reduccion horizontal por cara [%]", fr.get("reduccion_cara_horizontal_pct")),
        ])
        _headers(ws, r, ["Franja", "Z ini [m]", "Z fin [m]", "Shells",
                         "M22 eje [kN-m/m]", "M22 diseno [kN-m/m]",
                         "As vert [mm2/m]", "Gobierna vert",
                         "M11 eje pila", "M11 cara pila", "M11 vano",
                         "As horiz apoyo [mm2/m]", "As horiz vano [mm2/m]",
                         "V [kN/m]", "D/C V", "OK"])
        for i, f in enumerate(fr["franjas"], start=r + 1):
            ok = f["vertical"]["ok"] and f["horizontal_apoyo"]["ok"] \
                and f["horizontal_vano"]["ok"] and f["cortante"]["ok"]
            vals = [f["franja"], f["z_ini"], f["z_fin"], f["n_shells"],
                    f["M22"], f["M22_diseno"], f["vertical"]["As"], f["vertical"]["gobierna"],
                    f["M11_eje"], f["M11_cara"], f["M11_vano"],
                    f["horizontal_apoyo"]["As"], f["horizontal_vano"]["As"],
                    f["V"], f["cortante"]["ratio"], "SI" if ok else "NO"]
            for cix, v in enumerate(vals, start=1):
                ws.cell(row=i, column=cix, value=v)
            ws.cell(row=i, column=16).fill = OK_FILL if ok else BAD_FILL
        r += len(fr["franjas"]) + 2
        ws.cell(row=r, column=1, value="Nota").font = TITLE_FONT
        ws.cell(row=r + 1, column=1, value=fr.get("nota", "")).alignment = Alignment(wrap_text=True)
        _autofit(ws)

    # ---- Modal -----------------------------------------------------------
    mod = checks.get("modal") or {}
    if mod:
        ws = wb.create_sheet("Modal")
        r = _kv_sheet(ws, "Analisis modal", [
            ("Caso", mod["caso"]),
            ("Numero de modos", mod["n_modos"]),
            ("T1 [s]", mod["T1"]),
            ("Modo dominante en X", mod["modo_dominante_X"]),
            ("T del modo dominante en X [s]", mod["T_dominante_X"]),
            ("Modo dominante en Y", mod["modo_dominante_Y"]),
            ("T del modo dominante en Y [s]", mod["T_dominante_Y"]),
            ("Masa acumulada X", mod["masa_acumulada_X"]),
            ("Masa acumulada Y", mod["masa_acumulada_Y"]),
        ])
        _headers(ws, r, ["Modo", "T [s]", "f [Hz]", "Ux", "Uy", "Uz",
                         "Sum Ux", "Sum Uy", "Sum Uz"])
        for i, md in enumerate(mod["modos"], start=r + 1):
            for cix, v in enumerate([md["modo"], md["T"], md["f"], md["Ux"], md["Uy"],
                                     md["Uz"], md["SumUx"], md["SumUy"], md["SumUz"]], start=1):
                ws.cell(row=i, column=cix, value=v)
        _autofit(ws)

    # ---- Diseno de SAP ---------------------------------------------------
    ds = checks.get("diseno_sap") or {}
    if ds:
        ws = wb.create_sheet("DisenoSAP")
        r = _kv_sheet(ws, "Diseno de concreto corrido por SAP2000", [
            ("Codigo", ds.get("codigo")),
        ])
        if ds.get("viga_cabezal"):
            ws.cell(row=r, column=1, value="Viga cabezal").font = TITLE_FONT
            _headers(ws, r + 1, ["Frame", "X [m]", "As sup [mm2]", "As inf [mm2]",
                                 "Av/s [mm2/m]", "Al torsion [mm2]", "At/s torsion [mm2/m]",
                                 "Error", "Aviso"])
            for i, rec in enumerate(ds["viga_cabezal"], start=r + 2):
                for cix, v in enumerate([rec["frame"], rec["x"], rec["As_sup"], rec["As_inf"],
                                         rec["Av_s"], rec["Al_torsion"], rec["At_s_torsion"],
                                         rec["error"], rec["aviso"]], start=1):
                    ws.cell(row=i, column=cix, value=v)
            r += len(ds["viga_cabezal"]) + 3
        if ds.get("pilas"):
            ws.cell(row=r, column=1, value="Pilas").font = TITLE_FONT
            _headers(ws, r + 1, ["Frame", "X [m]", "As P-M-M [mm2]", "D/C P-M-M",
                                 "Av/s [mm2/m]", "Error", "Aviso"])
            for i, rec in enumerate(ds["pilas"], start=r + 2):
                for cix, v in enumerate([rec["frame"], rec["x"], rec["As_PMM"], rec["ratio_PMM"],
                                         rec["Av_s_mayor"], rec["error"], rec["aviso"]], start=1):
                    ws.cell(row=i, column=cix, value=v)
        _autofit(ws)

    # ---- Reacciones ------------------------------------------------------
    reac = checks.get("reaccion_vertical") or {}
    if reac:
        ws = wb.create_sheet("Reacciones")
        ws.cell(row=1, column=1, value="Reaccion vertical total por combinacion").font = TITLE_FONT
        _headers(ws, 2, ["Combinacion", "Suma Fz [kN]"])
        for i, (k, v) in enumerate(reac.items(), start=3):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=v)
        _autofit(ws)

    # ---- Despiece --------------------------------------------------------
    if schedule and schedule.get("marcas"):
        ws = wb.create_sheet("Despiece")
        _headers(ws, 1, ["Marca", "Elemento", "Posicion", "Diametro", "db [mm]",
                         "Forma", "Dims [mm]", "Largo [m]", "Cantidad",
                         "Largo total [m]", "Peso unit [kg]", "Peso total [kg]", "Nota"])
        for i, m in enumerate(schedule["marcas"], start=2):
            vals = [m["marca"], m["elemento"], m["posicion"], m["diametro"], m["db_mm"],
                    m["forma"], " x ".join("%.0f" % d for d in m["dims_mm"]),
                    m["largo_m"], m["cantidad"], m["largo_total_m"],
                    m["peso_unitario_kg"], m["peso_total_kg"], m["nota"]]
            for cix, v in enumerate(vals, start=1):
                ws.cell(row=i, column=cix, value=v)
        r = len(schedule["marcas"]) + 3
        ws.cell(row=r, column=1, value="Resumen por diametro").font = TITLE_FONT
        _headers(ws, r + 1, ["Diametro", "db [mm]", "Longitud [m]", "Peso [kg]", "Marcas"])
        for i, d in enumerate(schedule["por_diametro"], start=r + 2):
            for cix, v in enumerate([d["diametro"], d["db_mm"], d["largo_m"],
                                     d["peso_kg"], d["n_marcas"]], start=1):
                ws.cell(row=i, column=cix, value=v)
        r2 = r + 2 + len(schedule["por_diametro"]) + 1
        _kv_sheet(ws, "Totales", [
            ("Acero total [kg]", schedule["peso_total_kg"]),
            ("Concreto [m3]", schedule["volumen_concreto_m3"]),
            ("Cuantia [kg/m3]", schedule["cuantia_kg_m3"]),
        ] + [("Acero en %s [kg]" % k, v) for k, v in schedule["por_elemento"].items()]
          + [("", ""), ("Nota", schedule["nota"])], r2)
        _autofit(ws)

    # ---- Section cuts ----------------------------------------------------
    cuts = checks.get("section_cuts") or []
    if cuts:
        ws = wb.create_sheet("SectionCuts")
        _headers(ws, 1, ["Corte", "Caso", "F1 [kN]", "F2 [kN]", "F3 [kN]",
                         "M1 [kN-m]", "M2 [kN-m]", "M3 [kN-m]"])
        for i, rec in enumerate(cuts, start=2):
            for cix, v in enumerate([rec["cut"], rec["case"], rec["F1"], rec["F2"], rec["F3"],
                                     rec["M1"], rec["M2"], rec["M3"]], start=1):
                ws.cell(row=i, column=cix, value=round(v, 2) if isinstance(v, float) else v)
        _autofit(ws)

    return wb


def save_report(project: WallProject, report: dict, path: str,
                checks: dict | None = None, warnings: list[str] | None = None,
                schedule: dict | None = None) -> str:
    wb = build_report(project, report, checks, warnings, schedule)
    wb.save(path)
    return path
